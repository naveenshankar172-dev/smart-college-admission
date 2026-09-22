from pathlib import Path
from textwrap import wrap

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Inches, Pt, RGBColor
from pptx import Presentation
from pptx.dml.color import RGBColor as PptColor
from pptx.util import Inches as PptInches, Pt as PptPt
from PIL import Image, ImageDraw, ImageFont


ROOT = Path(__file__).resolve().parents[1]
OUT = ROOT / "docs" / "artifacts"
REPORT = ROOT / "docs" / "PROJECT_REPORT.md"


def font(size=24, bold=False):
    candidates = ["C:/Windows/Fonts/arial.ttf", "C:/Windows/Fonts/segoeui.ttf"]
    if bold:
        candidates = ["C:/Windows/Fonts/arialbd.ttf", "C:/Windows/Fonts/segoeuib.ttf"] + candidates
    for candidate in candidates:
        if Path(candidate).exists():
            return ImageFont.truetype(candidate, size)
    return ImageFont.load_default()


def make_diagram(name, title, nodes, links):
    width, height = 1800, 1050
    image = Image.new("RGB", (width, height), "#f7f4ee")
    draw = ImageDraw.Draw(image)
    draw.text((70, 45), title, fill="#17324d", font=font(44, True))
    positions = {}
    for index, (label, x, y, color) in enumerate(nodes):
        box = (x, y, x + 390, y + 120)
        positions[label] = box
        draw.rounded_rectangle(box, radius=18, fill=color, outline="#17324d", width=4)
        lines = wrap(label, 24)
        total = len(lines) * 30
        start = y + (120 - total) // 2
        for line_index, line in enumerate(lines):
            left = x + (390 - draw.textlength(line, font=font(25, True))) / 2
            draw.text((left, start + line_index * 30), line, fill="#17324d", font=font(25, True))
    for source, target in links:
        sx, sy, sw, sh = positions[source]
        tx, ty, tw, th = positions[target]
        start = (sx + sw // 2, sy + sh)
        end = (tx + tw // 2, ty)
        draw.line((start[0], start[1], end[0], end[1]), fill="#c55a3b", width=6)
        draw.polygon([(end[0] - 12, end[1] - 18), (end[0] + 12, end[1] - 18), end], fill="#c55a3b")
    image.save(OUT / name)


def create_diagrams():
    make_diagram("architecture.png", "System Architecture", [
        ("Student browser", 700, 150, "#d9e9f4"),
        ("React + TypeScript + Tailwind", 700, 340, "#c9e6d1"),
        ("Firebase Authentication / Google Sign-In", 700, 530, "#f4dfb5"),
        ("FastAPI REST API", 700, 720, "#f1c8b9"),
        ("SQLite + local uploads", 700, 910, "#d8c9e8"),
        ("UiPath -> OCR -> verification", 90, 720, "#f1c8b9"),
        ("Admin portal + notifications", 1310, 720, "#c9e6d1"),
    ], [("Student browser", "React + TypeScript + Tailwind"), ("React + TypeScript + Tailwind", "Firebase Authentication / Google Sign-In"), ("Firebase Authentication / Google Sign-In", "FastAPI REST API"), ("FastAPI REST API", "SQLite + local uploads"), ("UiPath -> OCR -> verification", "FastAPI REST API"), ("Admin portal + notifications", "FastAPI REST API")])
    make_diagram("er_diagram.png", "Entity Relationship Diagram", [
        ("Student", 700, 140, "#d9e9f4"), ("Applications", 150, 390, "#c9e6d1"), ("Documents", 700, 390, "#f4dfb5"), ("Scholarship Applications", 1250, 390, "#f1c8b9"), ("Scholarships", 1250, 700, "#d8c9e8"), ("Notifications", 150, 700, "#c9e6d1")
    ], [("Student", "Applications"), ("Student", "Documents"), ("Student", "Scholarship Applications"), ("Student", "Notifications"), ("Scholarship Applications", "Scholarships")])
    make_diagram("dfd_level_0.png", "DFD Level 0", [
        ("Student", 100, 450, "#d9e9f4"), ("Administrator", 1310, 450, "#d9e9f4"), ("Smart College Admission & Document Verification System", 600, 430, "#f1c8b9"), ("SQLite database", 700, 800, "#d8c9e8"), ("Firebase / UiPath / OCR", 650, 120, "#f4dfb5")
    ], [("Student", "Smart College Admission & Document Verification System"), ("Administrator", "Smart College Admission & Document Verification System"), ("Smart College Admission & Document Verification System", "SQLite database"), ("Firebase / UiPath / OCR", "Smart College Admission & Document Verification System")])
    make_diagram("use_case_diagram.png", "Use Case Overview", [
        ("Student", 80, 180, "#d9e9f4"), ("Sign in / onboarding", 500, 120, "#c9e6d1"), ("Admission / documents", 500, 340, "#c9e6d1"), ("Scholarships / eligibility", 500, 560, "#c9e6d1"), ("Notifications / profile", 500, 780, "#c9e6d1"), ("Administrator", 1310, 180, "#d9e9f4"), ("Review applications/documents", 1050, 340, "#f1c8b9"), ("Manage scholarships", 1050, 560, "#f1c8b9"), ("Reports / status updates", 1050, 780, "#f1c8b9")
    ], [("Student", "Sign in / onboarding"), ("Student", "Admission / documents"), ("Student", "Scholarships / eligibility"), ("Student", "Notifications / profile"), ("Administrator", "Review applications/documents"), ("Administrator", "Manage scholarships"), ("Administrator", "Reports / status updates")])
    make_diagram("activity_document_workflow.png", "Activity: Document Workflow", [
        ("Upload", 700, 120, "#d9e9f4"), ("Store metadata and file", 700, 300, "#c9e6d1"), ("Pending queue", 700, 480, "#f4dfb5"), ("UiPath calls process", 700, 660, "#f1c8b9"), ("OCR and consistency verification", 700, 840, "#d8c9e8")
    ], [("Upload", "Store metadata and file"), ("Store metadata and file", "Pending queue"), ("Pending queue", "UiPath calls process"), ("UiPath calls process", "OCR and consistency verification")])
    make_diagram("activity_authentication.png", "Activity: Student Authentication", [
        ("Google sign-in", 700, 120, "#d9e9f4"), ("Firebase identity", 700, 300, "#f4dfb5"), ("Existing student?", 700, 480, "#f1c8b9"), ("Onboarding", 220, 700, "#c9e6d1"), ("Student dashboard", 1180, 700, "#d8c9e8")
    ], [("Google sign-in", "Firebase identity"), ("Firebase identity", "Existing student?"), ("Existing student?", "Onboarding"), ("Existing student?", "Student dashboard"), ("Onboarding", "Student dashboard")])
    make_diagram("activity_admission.png", "Activity: Admission Workflow", [
        ("Login", 700, 120, "#d9e9f4"), ("Admission form", 700, 300, "#c9e6d1"), ("Submit", 700, 480, "#f4dfb5"), ("SQLite application", 700, 660, "#f1c8b9"), ("Admin review and notification", 700, 840, "#d8c9e8")
    ], [("Login", "Admission form"), ("Admission form", "Submit"), ("Submit", "SQLite application"), ("SQLite application", "Admin review and notification")])
    make_diagram("activity_scholarship.png", "Activity: Scholarship Workflow", [
        ("View scholarship", 700, 120, "#d9e9f4"), ("Evaluate configurable rules", 700, 300, "#c9e6d1"), ("Potentially eligible?", 700, 480, "#f4dfb5"), ("Apply", 220, 700, "#f1c8b9"), ("Review and status", 1180, 700, "#d8c9e8")
    ], [("View scholarship", "Evaluate configurable rules"), ("Evaluate configurable rules", "Potentially eligible?"), ("Potentially eligible?", "Apply"), ("Potentially eligible?", "Review and status"), ("Apply", "Review and status")])
    make_diagram("sequence_diagram.png", "Sequence: Document Processing", [
        ("Student / React", 100, 160, "#d9e9f4"), ("FastAPI", 700, 160, "#c9e6d1"), ("UiPath", 1300, 160, "#f4dfb5"), ("OCR + verification", 700, 560, "#f1c8b9"), ("SQLite", 1300, 560, "#d8c9e8")
    ], [("Student / React", "FastAPI"), ("FastAPI", "UiPath"), ("UiPath", "OCR + verification"), ("OCR + verification", "SQLite"), ("SQLite", "Student / React")])
    make_diagram("sequence_authentication.png", "Sequence: Google Authentication", [
        ("Student", 80, 160, "#d9e9f4"), ("React", 480, 160, "#c9e6d1"), ("Firebase", 880, 160, "#f4dfb5"), ("FastAPI / Admin", 1280, 160, "#f1c8b9"), ("SQLite", 880, 560, "#d8c9e8")
    ], [("Student", "React"), ("React", "Firebase"), ("Firebase", "React"), ("React", "FastAPI / Admin"), ("FastAPI / Admin", "SQLite")])
    make_diagram("sequence_scholarship.png", "Sequence: Scholarship Application", [
        ("Student", 80, 160, "#d9e9f4"), ("React", 480, 160, "#c9e6d1"), ("FastAPI", 880, 160, "#f1c8b9"), ("Eligibility", 1280, 160, "#f4dfb5"), ("SQLite + notification", 880, 560, "#d8c9e8")
    ], [("Student", "React"), ("React", "FastAPI"), ("FastAPI", "Eligibility"), ("Eligibility", "SQLite + notification"), ("SQLite + notification", "React")])
    make_diagram("sequence_admin_review.png", "Sequence: Admin Review", [
        ("Administrator", 80, 160, "#d9e9f4"), ("React", 480, 160, "#c9e6d1"), ("FastAPI", 880, 160, "#f1c8b9"), ("SQLite", 1280, 160, "#d8c9e8"), ("Student notification", 880, 560, "#f4dfb5")
    ], [("Administrator", "React"), ("React", "FastAPI"), ("FastAPI", "SQLite"), ("SQLite", "Student notification"), ("Student notification", "React")])
    make_diagram("dfd_level_1.png", "DFD Level 1", [
        ("Admission", 100, 160, "#d9e9f4"), ("Document verification", 100, 450, "#c9e6d1"), ("Scholarship", 100, 740, "#f4dfb5"), ("Administration", 1310, 450, "#f1c8b9"), ("FastAPI system processes", 650, 450, "#d8c9e8"), ("SQLite", 700, 850, "#d8c9e8")
    ], [("Admission", "FastAPI system processes"), ("Document verification", "FastAPI system processes"), ("Scholarship", "FastAPI system processes"), ("Administration", "FastAPI system processes"), ("FastAPI system processes", "SQLite")])


def add_docx_content(doc):
    for raw in REPORT.read_text(encoding="utf-8").splitlines():
        line = raw.strip()
        if not line:
            continue
        if line.startswith("# "):
            p = doc.add_heading(line[2:], 0)
        elif line.startswith("## "):
            p = doc.add_heading(line[3:], 1)
        elif line.startswith("### "):
            p = doc.add_heading(line[4:], 2)
        elif line.startswith("- ") or line.startswith("1. "):
            p = doc.add_paragraph(style="List Bullet")
            p.add_run(line[2:] if line.startswith("- ") else line[3:])
        elif line.startswith("|") or line.startswith(">") or line.startswith("`"):
            p = doc.add_paragraph(line)
        else:
            doc.add_paragraph(line)


def create_docx():
    doc = Document()
    section = doc.sections[0]
    section.top_margin = Inches(0.7)
    section.bottom_margin = Inches(0.7)
    styles = doc.styles
    styles["Normal"].font.name = "Aptos"
    styles["Normal"].font.size = Pt(10)
    styles["Heading 1"].font.color.rgb = RGBColor(23, 50, 77)
    styles["Heading 2"].font.color.rgb = RGBColor(197, 90, 59)
    add_docx_content(doc)
    footer = section.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
    footer.add_run("Smart College Admission & Document Verification System | Academic prototype")
    doc.save(OUT / "PROJECT_REPORT.docx")


SLIDES = [
    ("SMART COLLEGE ADMISSION & DOCUMENT VERIFICATION SYSTEM", "An integrated admission, document verification, scholarship screening and RPA automation platform\n[Team / student details]"),
    ("Problem Statement", "Manual document handling\nFragmented status tracking\nRepeated verification\nManual scholarship screening"),
    ("Objectives", "Secure identity path\nCentralized admission workflow\nOCR-assisted consistency checks\nConfigurable scholarship screening\nAdmin review and notifications"),
    ("Existing vs Proposed", "Conventional: forms, files, spreadsheets, repeated communication\nProposed: connected student, API, database, processing and review workflow"),
    ("System Architecture", "React + TypeScript -> Firebase identity -> FastAPI -> SQLite\nUiPath orchestrates pending documents; FastAPI performs OCR and verification"),
    ("Technology Stack", "React, TypeScript, Vite, Tailwind CSS\nPython, FastAPI, SQLAlchemy, Pydantic\nSQLite, Firebase, UiPath\nTesseract, pytesseract, PyMuPDF, pypdf, Pillow"),
    ("Authentication & Onboarding", "Google popup sign-in\nFirebase ID token\nFastAPI Admin verification\nFirst-time onboarding / returning session\nDevelopment mode is separately documented"),
    ("Admission Management", "Student profile\nApplication creation\nAdmin review and status changes\nRemarks and notifications"),
    ("Document Management", "PDF/JPG/JPEG/PNG upload\nMetadata and local storage\nPending processing state\nDownload/delete and duplicate-type controls"),
    ("OCR + Document Verification", "Text PDF -> pypdf\nScanned PDF/image -> PyMuPDF + Tesseract\nExtracted fields -> student comparison\nVerified / manual review / failed\nConsistency, not authenticity"),
    ("UiPath RPA Automation", "GET pending documents\nExtract document IDs\nPOST process request\nGET verification result\nStructured logging and errors\nOCR delegated to FastAPI"),
    ("Scholarship Eligibility & Application", "Minimum percentage\nCourse matching\nRequired documents\nIncome information flag\nPotentially Eligible / Requires Verification / Not Eligible"),
    ("Admin Dashboard & Review", "Live counts\nApplication review\nDocument verification queue\nScholarship CRUD\nScholarship application review\nRead-only settings"),
    ("Notifications & Status Tracking", "Admission events\nDocument upload and processing\nScholarship events\nMark one read / all read"),
    ("Database / ER Diagram", "Students -> applications, documents, scholarship applications, notifications\nScholarships -> scholarship applications\nSQLite is the prototype persistence layer"),
    ("Testing & Results", "Verified: Vite HTTP 200\nVerified: API health HTTP 200\nNot tested: production Firebase, Tesseract fixtures, UiPath execution, automated route suite\nNo fabricated performance metrics"),
    ("Limitations & Future Scope", "Current: development auth default, local SQLite/uploads, OCR sensitivity, manual review\nFuture: production hardening, PostgreSQL, cloud storage, audit logging, stronger authenticity analysis"),
    ("Conclusion", "A functional academic prototype integrating authentication, admission, documents, OCR, consistency verification, RPA orchestration, scholarships, administration and notifications\n[Thank you / questions]"),
]


def create_pptx():
    prs = Presentation()
    prs.slide_width = PptInches(13.333)
    prs.slide_height = PptInches(7.5)
    for index, (title, body) in enumerate(SLIDES):
        slide = prs.slides.add_slide(prs.slide_layouts[6])
        background = slide.background.fill
        background.solid()
        background.fore_color.rgb = PptColor(247, 244, 238)
        bar = slide.shapes.add_shape(1, 0, 0, PptInches(13.333), PptInches(0.22))
        bar.fill.solid(); bar.fill.fore_color.rgb = PptColor(197, 90, 59); bar.line.fill.background()
        text = slide.shapes.add_textbox(PptInches(0.85), PptInches(1.0), PptInches(11.7), PptInches(1.2))
        tf = text.text_frame; tf.clear()
        p = tf.paragraphs[0]; p.text = title; p.font.size = PptPt(28 if index else 31); p.font.bold = True; p.font.color.rgb = PptColor(23, 50, 77)
        body_box = slide.shapes.add_textbox(PptInches(1.0), PptInches(2.45), PptInches(11.2), PptInches(3.8))
        tf = body_box.text_frame; tf.clear(); tf.word_wrap = True
        for line_index, line in enumerate(body.split("\n")):
            p = tf.paragraphs[0] if line_index == 0 else tf.add_paragraph()
            p.text = line; p.font.size = PptPt(22); p.font.color.rgb = PptColor(45, 58, 66); p.space_after = PptPt(16)
        footer = slide.shapes.add_textbox(PptInches(11.9), PptInches(7.02), PptInches(0.7), PptInches(0.25))
        footer.text_frame.paragraphs[0].text = str(index + 1)
        footer.text_frame.paragraphs[0].font.size = PptPt(11)
    prs.save(OUT / "PROJECT_PRESENTATION.pptx")


def main():
    OUT.mkdir(parents=True, exist_ok=True)
    create_diagrams()
    create_docx()
    create_pptx()
    print(f"Generated artifacts in {OUT}")


if __name__ == "__main__":
    main()